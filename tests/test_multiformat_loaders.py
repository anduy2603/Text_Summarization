from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from app.services.input import InputLoadError, process_from_bytes, process_from_text
from app.services.input.loaders.docx_loader import load_docx_bytes
from app.services.input.loaders.pdf_loader import load_pdf_bytes
from app.services.input.loaders.txt_loader import load_txt_bytes
from scripts.shared.multiformat_experiment import (
    build_docx_bytes,
    build_pdf_bytes,
    build_txt_bytes,
    compute_ingestion_fidelity,
)


VIET_SAMPLE = (
    "Thành phố Hồ Chí Minh triển khai tuyến xe buýt điện. "
    "Người dân phản hồi tích cực về chất lượng dịch vụ."
)


def test_load_txt_utf8() -> None:
    raw = build_txt_bytes(VIET_SAMPLE)
    text = load_txt_bytes(raw)
    assert "Hồ Chí Minh" in text


def test_load_txt_utf8_bom() -> None:
    raw = b"\xef\xbb\xbf" + VIET_SAMPLE.encode("utf-8")
    text = load_txt_bytes(raw)
    assert "Hồ Chí Minh" in text


def test_load_txt_windows_1258() -> None:
    # ASCII-safe payload encodable in cp1258; chardet should still decode via loader.
    raw = "Tai lieu tom tat dia phuong.".encode("cp1258")
    text = load_txt_bytes(raw)
    assert "tom tat" in text


def test_load_docx_paragraph_and_table() -> None:
    buf = BytesIO()
    doc = Document()
    doc.add_paragraph(VIET_SAMPLE)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Cột A"
    table.rows[0].cells[1].text = "Cột B"
    doc.save(buf)
    text = load_docx_bytes(buf.getvalue())
    assert "Hồ Chí Minh" in text
    assert "Cột A" in text


def test_load_pdf_single_page() -> None:
    raw = build_pdf_bytes(VIET_SAMPLE)
    text = load_pdf_bytes(raw)
    assert len(text.strip()) > 20


def test_load_pdf_blank_raises_clear_message() -> None:
    import fitz

    doc = fitz.open()
    doc.new_page()
    raw = doc.tobytes()
    doc.close()
    with pytest.raises(InputLoadError, match="không chứa văn bản"):
        load_pdf_bytes(raw)


def test_process_from_bytes_txt_matches_text_baseline() -> None:
    baseline = process_from_text(VIET_SAMPLE)
    from_txt = process_from_bytes("sample.txt", build_txt_bytes(VIET_SAMPLE))
    fidelity = compute_ingestion_fidelity(baseline, from_txt)
    assert fidelity["ingest_exact_match"] is True
    assert fidelity["ingest_rougeL_f"] == pytest.approx(1.0)


def test_process_from_bytes_docx_matches_text_baseline() -> None:
    baseline = process_from_text(VIET_SAMPLE)
    from_docx = process_from_bytes("sample.docx", build_docx_bytes(VIET_SAMPLE))
    fidelity = compute_ingestion_fidelity(baseline, from_docx)
    assert fidelity["ingest_exact_match"] is True


def test_process_from_bytes_pdf_non_empty() -> None:
    processed = process_from_bytes("sample.pdf", build_pdf_bytes(VIET_SAMPLE))
    assert processed.source_type == "pdf"
    assert len(processed.cleaned_text) > 0
    assert len(processed.sentences) >= 1


def test_load_txt_empty_raises() -> None:
    with pytest.raises(InputLoadError, match="empty"):
        load_txt_bytes(b"")
