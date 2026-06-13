from __future__ import annotations

import asyncio
import io
import re
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import Response

from app.core.config import settings
from app.schemas.common import ExportDocxRequest, SummarizeRequest, SummarizeResponse, SummaryControls
from app.schemas.input import UrlIngestRequest
from app.services.input import (
    InputLoadError,
    InputValidationError,
    process_from_bytes,
    process_from_text,
    process_from_url,
)
from app.services.summarization import (
    list_planned_summary_engines,
    list_supported_summary_engines,
    summarize_processed_input,
)
from app.services.summarization.summary_service import SummaryEngineNotReadyError, UnsupportedSummaryEngineError
from app.core.rate_limit import limiter

router = APIRouter()


def _map_input_errors(exc: Exception) -> HTTPException:
    if isinstance(exc, InputValidationError):
        return HTTPException(status_code=400, detail=str(exc))
    if isinstance(exc, InputLoadError):
        return HTTPException(status_code=422, detail=str(exc))
    return HTTPException(status_code=500, detail="Unexpected input error.")


def _summary_controls_from_query(
    max_sentences: Optional[int] = Query(
        default=None,
        ge=1,
        le=20,
        description="Preferred top-k sentence count. Takes priority over ratio when provided.",
    ),
    ratio: Optional[float] = Query(
        default=None,
        gt=0.0,
        le=1.0,
        description="Optional extractive ratio. Used only when max_sentences is None.",
    ),
    engine: Optional[str] = Query(
        default=None,
        description="Optional summarization engine override.",
    ),
) -> SummaryControls:
    return SummaryControls(max_sentences=max_sentences, ratio=ratio, engine=engine)


@router.get("/engines")
async def list_engines() -> dict[str, object]:
    """
    Capability endpoint for frontend/runtime discovery.
    Keep this lightweight so clients can refresh available engines at startup.
    """
    supported = list_supported_summary_engines()
    planned = list_planned_summary_engines()
    preferred = settings.summary_engine.strip().lower()
    if preferred in supported:
        default_engine: str | None = preferred
    else:
        default_engine = supported[0] if supported else None
    return {
        "supported_engines": supported,
        "planned_engines": planned,
        "default_engine": default_engine,
        "default_max_sentences": settings.summary_max_sentences,
    }


@router.post("/summarize", response_model=SummarizeResponse)
@limiter.limit("10/minute")
async def summarize(request: Request, payload: SummarizeRequest) -> SummarizeResponse:
    try:
        processed = await asyncio.to_thread(process_from_text, payload.text)
    except (InputValidationError, InputLoadError) as exc:
        raise _map_input_errors(exc) from exc
    try:
        return await asyncio.to_thread(
            summarize_processed_input,
            processed,
            payload.max_sentences,
            payload.ratio,
            payload.engine,
        )
    except SummaryEngineNotReadyError as exc:
        raise HTTPException(status_code=501, detail=str(exc)) from exc
    except UnsupportedSummaryEngineError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.post("/summarize/file", response_model=SummarizeResponse)
@limiter.limit("10/minute")
async def summarize_file(
    request: Request,
    file: UploadFile = File(...),
    controls: SummaryControls = Depends(_summary_controls_from_query),
) -> SummarizeResponse:
    content = await file.read()
    try:
        processed = await asyncio.to_thread(process_from_bytes, file.filename or "", content)
    except (InputValidationError, InputLoadError) as exc:
        raise _map_input_errors(exc) from exc
    try:
        return await asyncio.to_thread(
            summarize_processed_input,
            processed,
            controls.max_sentences,
            controls.ratio,
            controls.engine,
        )
    except SummaryEngineNotReadyError as exc:
        raise HTTPException(status_code=501, detail=str(exc)) from exc
    except UnsupportedSummaryEngineError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.post("/summarize/url", response_model=SummarizeResponse)
@limiter.limit("10/minute")
async def summarize_url(
    request: Request,
    payload: UrlIngestRequest,
    controls: SummaryControls = Depends(_summary_controls_from_query),
) -> SummarizeResponse:
    try:
        processed = await asyncio.to_thread(process_from_url, payload.url)
    except (InputValidationError, InputLoadError) as exc:
        raise _map_input_errors(exc) from exc
    try:
        return await asyncio.to_thread(
            summarize_processed_input,
            processed,
            controls.max_sentences,
            controls.ratio,
            controls.engine,
        )
    except SummaryEngineNotReadyError as exc:
        raise HTTPException(status_code=501, detail=str(exc)) from exc
    except UnsupportedSummaryEngineError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.post("/export/docx")
async def export_summary_as_docx(payload: ExportDocxRequest) -> Response:
    """Export a summary as a formatted DOCX file."""
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise HTTPException(status_code=501, detail="python-docx not installed.") from exc

    doc = Document()

    heading = doc.add_heading(payload.title or "Tóm tắt", level=1)
    heading.runs[0].font.size = Pt(16)

    lines = [ln.strip() for ln in payload.summary.splitlines() if ln.strip()]
    for line in lines:
        clean = line.lstrip("•·-* ")
        if clean:
            para = doc.add_paragraph(clean)
            para.runs[0].font.size = Pt(12) if para.runs else None

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    from urllib.parse import quote as urlquote

    raw_name = (payload.title or "tom-tat")[:60]
    # ASCII fallback filename (latin-1 safe)
    ascii_name = re.sub(r"[^\x20-\x7e]", "_", raw_name).strip("_") or "tom-tat"
    ascii_name = re.sub(r"[^\w\-. ]", "_", ascii_name).replace(" ", "_")
    # RFC 5987 UTF-8 encoded filename for modern clients
    utf8_name = urlquote(raw_name + ".docx", safe="")

    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{ascii_name}.docx"; '
                f"filename*=UTF-8''{utf8_name}"
            )
        },
    )
